"""
core/reports.py - Geração de relatórios Excel (openpyxl) e PDF (fpdf2).
"""
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from fpdf import FPDF
from utils.helpers import get_logger, ensure_dir, format_cnpj

logger = get_logger(__name__)

# Paleta de cores para status
COLOR_OK = "C6EFCE"       # Verde claro
COLOR_ALERTA = "FFEB9C"   # Amarelo claro
COLOR_CRITICO = "FFC7CE"  # Vermelho claro
COLOR_REDUCAO = "D6EAF8"  # Azul claro
COLOR_HEADER = "2F5496"   # Azul escuro (cabeçalho)
COLOR_WHITE = "FFFFFF"

FILL_OK = PatternFill("solid", fgColor=COLOR_OK)
FILL_ALERTA = PatternFill("solid", fgColor=COLOR_ALERTA)
FILL_CRITICO = PatternFill("solid", fgColor=COLOR_CRITICO)
FILL_REDUCAO = PatternFill("solid", fgColor=COLOR_REDUCAO)
FILL_HEADER = PatternFill("solid", fgColor=COLOR_HEADER)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

HEADERS = [
    "Arquivo", "N. Nota", "CNPJ Fornecedor", "Posto Referência", "Descrição",
    "Qtd", "Valor Total (R$)", "Preço XML (R$)", "Preço API (R$)",
    "Dif. (R$)", "Dif. %", "Status", "Data Processamento"
]

FIELD_MAP = [
    "nome_arquivo", "numero_nota", "cnpj_fornecedor", "posto_referencia", "descricao",
    "qtd", "valor_total", "preco_xml", "preco_api",
    "diff_abs", "diff_pct", "status", "data_processamento"
]


def export_excel(rows: list, output_path: str) -> bool:
    """
    Gera relatório Excel com formatação condicional de cores por status.

    Parâmetros:
        rows: lista de dicts com os resultados
        output_path: caminho completo do arquivo .xlsx de saída

    Retorna True em caso de sucesso.
    """
    try:
        ensure_dir(os.path.dirname(output_path))
        wb = Workbook()
        ws = wb.active
        ws.title = "Comparação de Preços"

        # Linha de título
        ws.merge_cells("A1:M1")
        title_cell = ws["A1"]
        title_cell.value = f"Relatório de Comparação de Preços — Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        title_cell.font = Font(bold=True, color=COLOR_WHITE, size=12)
        title_cell.fill = FILL_HEADER
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        # Cabeçalhos
        for col_idx, header in enumerate(HEADERS, start=1):
            cell = ws.cell(row=2, column=col_idx, value=header)
            cell.font = Font(bold=True, color=COLOR_WHITE)
            cell.fill = FILL_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = THIN_BORDER
        ws.row_dimensions[2].height = 30

        # Dados
        for row_idx, row in enumerate(rows, start=3):
            status = row.get("status", "OK")
            fill = (
                FILL_OK if status == "OK"
                else FILL_ALERTA if status == "ALERTA"
                else FILL_CRITICO if status == "CRITICO"
                else FILL_REDUCAO
            )

            for col_idx, field in enumerate(FIELD_MAP, start=1):
                value = row.get(field, "")
                if field == "cnpj_fornecedor":
                    value = format_cnpj(value)
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = fill
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center")

                # Formatação numérica
                if field in ("valor_total", "preco_xml", "preco_api", "diff_abs"):
                    cell.number_format = '#,##0.0000'
                elif field == "diff_pct":
                    cell.number_format = '0.00"%"'
                elif field == "qtd":
                    cell.number_format = '#,##0.0000'

        # Larguras de coluna
        col_widths = [30, 12, 20, 28, 35, 10, 18, 18, 18, 18, 12, 12, 22]
        for col_idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Congelar painel no cabeçalho
        ws.freeze_panes = "A3"

        # Aba de resumo
        ws_sum = wb.create_sheet("Resumo")
        _write_summary_sheet(ws_sum, rows)

        wb.save(output_path)
        logger.info("Excel exportado: %s (%d linhas)", output_path, len(rows))
        return True

    except Exception as e:
        logger.error("Erro ao gerar Excel: %s", e)
        return False


def _write_summary_sheet(ws, rows: list):
    """Escreve aba de resumo com contagem por status."""
    from collections import Counter
    counts = Counter(r.get("status", "OK") for r in rows)

    ws["A1"] = "Resumo por Status"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Status"
    ws["B2"] = "Quantidade"
    for cell in [ws["A2"], ws["B2"]]:
        cell.font = Font(bold=True, color=COLOR_WHITE)
        cell.fill = FILL_HEADER

    for row_idx, (status, qty) in enumerate(
        [("OK", counts.get("OK", 0)),
         ("ALERTA", counts.get("ALERTA", 0)),
         ("CRITICO", counts.get("CRITICO", 0)),
         ("REDUCAO", counts.get("REDUCAO", 0))],
        start=3
    ):
        fill = (
            FILL_OK if status == "OK"
            else FILL_ALERTA if status == "ALERTA"
            else FILL_CRITICO if status == "CRITICO"
            else FILL_REDUCAO
        )
        c_status = ws.cell(row=row_idx, column=1, value=status)
        c_qty = ws.cell(row=row_idx, column=2, value=qty)
        c_status.fill = fill
        c_qty.fill = fill

    ws["A6"] = "Total"
    ws["B6"] = len(rows)
    ws["A6"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 15


def export_pdf(rows: list, output_path: str) -> bool:
    """
    Gera relatório PDF com cabeçalho, data e tabela de resultados.

    Parâmetros:
        rows: lista de dicts com os resultados
        output_path: caminho completo do arquivo .pdf de saída

    Retorna True em caso de sucesso.
    """
    try:
        ensure_dir(os.path.dirname(output_path))
        pdf = FPDF(orientation="L", unit="mm", format="A4")
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_margins(10, 10, 10)

        # Cabeçalho
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Relatório de Comparação de Preços", align="C", ln=True)
        pdf.set_font("Helvetica", "", 9)
        pdf.cell(0, 6, f"Gerado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", align="C", ln=True)
        pdf.ln(4)

        # Resumo
        from collections import Counter
        counts = Counter(r.get("status", "OK") for r in rows)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 6, f"Total: {len(rows)}  |  OK: {counts.get('OK',0)}  |  ALERTA: {counts.get('ALERTA',0)}  |  CRÍTICO: {counts.get('CRITICO',0)}", ln=True)
        pdf.ln(3)

        # Cabeçalho da tabela
        col_defs = [
            ("Arquivo", 28),
            ("N. Nota", 14),
            ("CNPJ", 28),
            ("Posto Ref.", 34),
            ("Descrição", 32),
            ("Qtd", 14),
            ("V.Total", 18),
            ("P.XML", 18),
            ("P.API", 18),
            ("Dif.", 16),
            ("Dif.%", 14),
            ("Status", 14),
        ]

        pdf.set_font("Helvetica", "B", 7)
        pdf.set_fill_color(47, 84, 150)
        pdf.set_text_color(255, 255, 255)
        for label, width in col_defs:
            pdf.cell(width, 7, label, border=1, align="C", fill=True)
        pdf.ln()

        # Linhas de dados
        pdf.set_font("Helvetica", "", 6.5)
        pdf.set_text_color(0, 0, 0)

        for row in rows:
            status = row.get("status", "OK")
            if status == "OK":
                pdf.set_fill_color(198, 239, 206)
            elif status == "ALERTA":
                pdf.set_fill_color(255, 235, 156)
            elif status == "REDUCAO":
                pdf.set_fill_color(214, 234, 248)
            else:
                pdf.set_fill_color(255, 199, 206)

            values = [
                _truncate(str(row.get("nome_arquivo", "")), 20),
                str(row.get("numero_nota", "")),
                format_cnpj(row.get("cnpj_fornecedor", "")),
                _truncate(str(row.get("posto_referencia", "")), 26),
                _truncate(str(row.get("descricao", "")), 24),
                f"{row.get('qtd', 0):.2f}",
                f"{row.get('valor_total', 0):.4f}",
                f"{row.get('preco_xml', 0):.4f}",
                f"{row.get('preco_api', 0):.4f}",
                f"{row.get('diff_abs', 0):.4f}",
                f"{row.get('diff_pct', 0):.2f}%",
                status,
            ]

            for (_, width), value in zip(col_defs, values):
                pdf.cell(width, 6, value, border=1, fill=True)
            pdf.ln()

        pdf.output(output_path)
        logger.info("PDF exportado: %s (%d linhas)", output_path, len(rows))
        return True

    except Exception as e:
        logger.error("Erro ao gerar PDF: %s", e)
        return False


def _truncate(text: str, max_len: int) -> str:
    """Trunca texto para caber na célula do PDF."""
    if len(text) <= max_len:
        return text
    if max_len <= 3:
        return text[:max_len]
    return text[:max_len - 3] + "..."
